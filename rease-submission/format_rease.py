import os
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def format_document(file_path):
    print(f"Loading document: {file_path}")
    doc = Document(file_path)

    # 1. Set A4 Page Size and Margins (Left/Top: 3.0 cm, Right/Bottom: 2.0 cm)
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(3.0)
        section.left_margin = Cm(3.0)
        section.bottom_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    in_abstract_or_resumo = False

    # 2. Iterate paragraphs and format
    for i, paragraph in enumerate(doc.paragraphs):
        style_name = paragraph.style.name.lower()
        text = paragraph.text.strip()
        
        # Identify title paragraph
        is_title = style_name == 'title' or (i == 0 and len(text) > 0)
        
        if is_title:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(12)
            
            # Style runs
            for run in paragraph.runs:
                run.bold = True
                run.font.name = 'Times New Roman'
                run.font.size = Pt(14)  # Title slightly larger
                run.text = run.text.upper()
            continue

        # Handle headings (like "1. Introduction", "Abstract", etc.)
        if paragraph.style.name.startswith('Heading'):
            # Convert text of headings to upper case
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            
            # Check if this heading starts abstract or resumo
            heading_text = text.upper()
            if "ABSTRACT" in heading_text or "RESUMO" in heading_text:
                in_abstract_or_resumo = True
            else:
                in_abstract_or_resumo = False

            for run in paragraph.runs:
                run.bold = True
                run.font.name = 'Times New Roman'
                run.font.size = Pt(12)
                run.text = run.text.upper()
            continue

        # Check if the paragraph text indicates abstract keywords or is part of abstract
        if text.upper().startswith("KEYWORDS:") or text.upper().startswith("PALAVRAS-CHAVE:"):
            # Keywords paragraph
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.15
            paragraph.paragraph_format.space_after = Pt(12)
            for run in paragraph.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(11)
            continue

        # Handle normal paragraphs
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        if in_abstract_or_resumo:
            # Abstract/Resumo text formatting (no first-line indent, spacing 1.15)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.15
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(6)
            for run in paragraph.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(12)
        else:
            # Main body normal paragraphs (first-line indent 1.25 cm, spacing 1.5)
            # Skip empty paragraphs
            if len(text) > 0:
                paragraph.paragraph_format.first_line_indent = Cm(1.25)
                paragraph.paragraph_format.line_spacing = 1.5
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(6)
            else:
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                
            for run in paragraph.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(12)

    # 3. Format Tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    paragraph.paragraph_format.first_line_indent = Cm(0)
                    paragraph.paragraph_format.line_spacing = 1.15
                    paragraph.paragraph_format.space_before = Pt(2)
                    paragraph.paragraph_format.space_after = Pt(2)
                    for run in paragraph.runs:
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(10) # 10 pt for table content

    print("Saving formatted document...")
    doc.save(file_path)
    print("Document formatted successfully!")

if __name__ == "__main__":
    import sys
    target = sys.argv[1] if len(sys.argv) > 1 else "rease_article.docx"
    format_document(target)
